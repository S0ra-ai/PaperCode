"""导出服务 — 将 Markdown 转换为 DOCX / PDF"""

from __future__ import annotations

import io
import re
import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)


def markdown_to_docx(markdown_text: str) -> bytes:
    """将 Markdown 报告转为 DOCX 字节流"""
    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "SimSun"
    font.size = Pt(12)

    lines = markdown_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        # Headings
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            text = re.sub(r"^\d+\.\s", "", line).strip()
            doc.add_paragraph(text, style="List Number")
        elif line.strip() == "":
            pass
        elif line.startswith("```"):
            # 代码块：收集到闭合 ```
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            code_text = "\n".join(code_lines)
            p = doc.add_paragraph()
            run = p.add_run(code_text)
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        else:
            # 普通段落 — 去除行内 Markdown
            text = line.strip()
            text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
            text = re.sub(r"\*(.+?)\*", r"\1", text)
            text = re.sub(r"`(.+?)`", r"\1", text)
            doc.add_paragraph(text)

        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF 导出（fpdf2）
# ---------------------------------------------------------------------------

_FONT_SEARCH_PATHS = [
    Path("C:/Windows/Fonts"),
    Path("/usr/share/fonts/truetype"),
    Path("/usr/share/fonts"),
    Path("/System/Library/Fonts"),
]


def _find_font(names: list[str]) -> str | None:
    """在系统路径中查找 .ttf 字体文件"""
    for base in _FONT_SEARCH_PATHS:
        if not base.exists():
            continue
        for name in names:
            candidate = base / name
            if candidate.exists():
                return str(candidate)
            # 递归查找
            for f in base.rglob(name):
                return str(f)
    return None


def markdown_to_pdf(markdown_text: str) -> bytes:
    """将 Markdown 转为支持中文的 PDF 字节流"""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    # 注册中文字体
    font_path = _find_font(["simhei.ttf", "SimHei.ttf", "msyh.ttc", "NotoSansSC-Regular.ttf"])
    if font_path:
        try:
            pdf.add_font("zh", "", font_path, uni=True)
            pdf.add_font("zh", "B", font_path, uni=True)
            zh_font = "zh"
        except Exception as e:
            logger.warning("加载中文字体失败 %s: %s", font_path, e)
            zh_font = "Helvetica"
    else:
        logger.warning("未找到中文字体，PDF 可能无法渲染 CJK 文字")
        zh_font = "Helvetica"

    pdf.set_font(zh_font, "", 11)
    effective_width = pdf.w - pdf.l_margin - pdf.r_margin

    lines = markdown_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith("# "):
            pdf.set_font(zh_font, "B", 18)
            pdf.ln(6)
            pdf.multi_cell(effective_width, 9, line[2:].strip())
            pdf.ln(4)
            pdf.set_font(zh_font, "", 11)
        elif line.startswith("## "):
            pdf.set_font(zh_font, "B", 15)
            pdf.ln(5)
            pdf.multi_cell(effective_width, 8, line[3:].strip())
            pdf.ln(3)
            pdf.set_font(zh_font, "", 11)
        elif line.startswith("### "):
            pdf.set_font(zh_font, "B", 13)
            pdf.ln(3)
            pdf.multi_cell(effective_width, 7, line[4:].strip())
            pdf.ln(2)
            pdf.set_font(zh_font, "", 11)
        elif line.startswith("#### "):
            pdf.set_font(zh_font, "B", 12)
            pdf.ln(2)
            pdf.multi_cell(effective_width, 7, line[5:].strip())
            pdf.ln(2)
            pdf.set_font(zh_font, "", 11)
        elif line.startswith("- ") or line.startswith("* "):
            text = "  ·  " + line[2:].strip()
            pdf.multi_cell(effective_width, 6, _strip_inline_md(text))
        elif re.match(r"^\d+\.\s", line):
            pdf.multi_cell(effective_width, 6, _strip_inline_md(line.strip()))
        elif line.strip() == "":
            pdf.ln(3)
        elif line.startswith("```"):
            # 代码块 — 使用中文字体以支持代码中的 CJK 字符
            pdf.set_font(zh_font, "", 9)
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code_lines.append(lines[i])
                i += 1
            for cl in code_lines:
                pdf.multi_cell(effective_width, 5, cl)
            pdf.set_font(zh_font, "", 11)
            pdf.ln(2)
        else:
            text = _strip_inline_md(line.strip())
            if text:
                pdf.multi_cell(effective_width, 6, text)

        i += 1

    return pdf.output()


def _strip_inline_md(text: str) -> str:
    """去除行内 Markdown 格式标记"""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text


# ---------------------------------------------------------------------------
# LaTeX 导出
# ---------------------------------------------------------------------------

def markdown_to_latex_zip(
    markdown_text: str,
    title: str = "",
    author: str = "",
    citations: list[dict] | None = None,
    citation_style: str = "GB/T 7714",
) -> bytes:
    """将 Markdown 转为 ZIP 包（含 main.tex + references.bib）"""
    from intelligen_engine.pipeline.export.latex_exporter import (
        markdown_to_latex_zip as _impl,
    )
    return _impl(
        markdown_text,
        title=title,
        author=author,
        citations=citations,
        citation_style=citation_style,
    )
