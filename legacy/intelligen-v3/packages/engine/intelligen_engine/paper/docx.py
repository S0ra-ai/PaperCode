"""DOCX rendering for paper-mode drafts."""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from intelligen_engine.paper.models import FormatSpec, PaperDraft


def render_docx(draft: PaperDraft, spec: FormatSpec) -> bytes:
    doc = Document()
    _configure_document(doc, spec)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(draft.title)
    title_run.bold = True
    title_run.font.size = Pt(16)

    authors = doc.add_paragraph(", ".join(draft.authors))
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(draft.abstract)

    doc.add_paragraph("Keywords: " + "; ".join(draft.keywords))

    for section in draft.sections:
        doc.add_heading(section.title, level=1)
        for paragraph in section.paragraphs:
            doc.add_paragraph(paragraph)

    if draft.bibliography.strip():
        doc.add_heading("References", level=1)
        doc.add_paragraph(draft.bibliography.strip())

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _configure_document(doc: Document, spec: FormatSpec) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10 if spec.venue_label.upper() == "IEEE" else 12)
