"""Structured MCP-style research tools.

The functions in this module are deliberately plain Python callables so tests and
the TypeScript runtime bridge can use them without requiring a live MCP server.
`mcp_server.py` adapts the same registry to the MCP SDK when installed.
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Callable

from .crawler import crawl_url, scrapling_status
from .external_tools import discover_external_tools
from .models import ToolArtifact, ToolEvent, ToolResult


ToolHandler = Callable[[dict[str, Any]], ToolResult]


def tool_diagnostics(input_data: dict[str, Any] | None = None) -> ToolResult:
    input_data = input_data or {}
    report = discover_external_tools(input_data.get("config_path"))
    return ToolResult.ok(
        "tool_diagnostics",
        artifacts=[
            ToolArtifact(
                title="External tool diagnostics",
                data=report.model_dump(),
            )
        ],
        events=[
            ToolEvent(
                type="python.tools.diagnostics",
                payload={
                    "available": [name for name, item in report.tools.items() if item.status == "available"],
                    "missing": [name for name, item in report.tools.items() if item.status != "available"],
                },
            )
        ],
        metadata=report.model_dump(),
    )


def pdf_probe(input_data: dict[str, Any]) -> ToolResult:
    file = Path(str(input_data["path"])).expanduser()
    if not file.exists():
        return ToolResult.fail("pdf_probe", f"PDF not found: {file}")
    metadata: dict[str, Any] = {
        "path": str(file),
        "size_bytes": file.stat().st_size,
        "looks_like_pdf": file.read_bytes()[:5] == b"%PDF-",
    }
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(str(file))
        metadata["pages"] = len(reader.pages)
    except Exception as error:  # pragma: no cover - optional dependency/runtime variability
        metadata["pages"] = None
        metadata["parser_warning"] = str(error)
    return ToolResult.ok("pdf_probe", metadata=metadata, artifacts=[ToolArtifact(title="PDF probe", data=metadata)])


def dataset_metadata_probe(input_data: dict[str, Any]) -> ToolResult:
    file = Path(str(input_data["path"])).expanduser()
    if not file.exists():
        return ToolResult.fail("dataset_metadata_probe", f"Dataset not found: {file}")
    suffix = file.suffix.lower()
    metadata: dict[str, Any] = {"path": str(file), "size_bytes": file.stat().st_size, "format": suffix.lstrip(".")}
    if suffix == ".csv":
        with file.open("r", encoding=input_data.get("encoding", "utf-8"), newline="") as handle:
            reader = csv.reader(handle)
            header = next(reader, [])
            rows = sum(1 for _ in reader)
        metadata.update({"columns": header, "rows": rows})
    elif suffix == ".json":
        data = json.loads(file.read_text(encoding=input_data.get("encoding", "utf-8")))
        metadata.update({"json_type": type(data).__name__})
        if isinstance(data, dict):
            metadata["keys"] = sorted(data.keys())
        if isinstance(data, list):
            metadata["items"] = len(data)
    elif suffix == ".jsonl":
        lines = [line for line in file.read_text(encoding=input_data.get("encoding", "utf-8")).splitlines() if line.strip()]
        metadata.update({"items": len(lines)})
    else:
        metadata["warning"] = "Unsupported dataset format; only size and suffix were recorded."
    return ToolResult.ok(
        "dataset_metadata_probe",
        artifacts=[ToolArtifact(title="Dataset metadata", data=metadata)],
        events=[ToolEvent(type="python.dataset.probed", payload=metadata)],
        metadata=metadata,
    )


METRIC_PATTERN = re.compile(r"\b(?P<name>loss|accuracy|acc|f1|precision|recall)\b\s*[:=]\s*(?P<value>-?\d+(?:\.\d+)?)", re.I)


def experiment_log_parse(input_data: dict[str, Any]) -> ToolResult:
    file = Path(str(input_data["path"])).expanduser()
    if not file.exists():
        return ToolResult.fail("experiment_log_parse", f"Experiment log not found: {file}")
    metrics: list[dict[str, Any]] = []
    for line_number, line in enumerate(file.read_text(encoding=input_data.get("encoding", "utf-8")).splitlines(), start=1):
        for match in METRIC_PATTERN.finditer(line):
            metrics.append(
                {
                    "line": line_number,
                    "name": match.group("name").lower(),
                    "value": float(match.group("value")),
                }
            )
    metadata = {"path": str(file), "metrics": metrics, "metric_count": len(metrics)}
    return ToolResult.ok(
        "experiment_log_parse",
        artifacts=[ToolArtifact(title="Experiment metrics", data=metadata)],
        events=[ToolEvent(type="python.experiment.parsed", payload={"metric_count": len(metrics)})],
        metadata=metadata,
    )


def web_crawl_probe(input_data: dict[str, Any]) -> ToolResult:
    record = crawl_url(input_data)
    status = "succeeded" if record.get("status") == "fetched" else "skipped"
    metadata = {
        "scrapling": scrapling_status(),
        "record": record,
    }
    if status == "skipped":
        return ToolResult(
            tool_name="web_crawl_probe",
            status="skipped",
            artifacts=[ToolArtifact(title="Crawler source record", data=record)],
            events=[ToolEvent(type="python.crawler.skipped", payload=metadata)],
            metadata=metadata,
            error="; ".join(str(item) for item in record.get("diagnostics", [])) or None,
        )
    return ToolResult.ok(
        "web_crawl_probe",
        artifacts=[ToolArtifact(title="Crawler source record", data=record)],
        events=[ToolEvent(type="python.crawler.fetched", payload=metadata)],
        metadata=metadata,
    )


def docx_generation(input_data: dict[str, Any]) -> ToolResult:
    output = Path(str(input_data["output_path"])).expanduser()
    title = str(input_data.get("title", "paper_fucker draft"))
    paragraphs = [str(item) for item in input_data.get("paragraphs", [])]
    try:
        from docx import Document  # type: ignore

        output.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading(title, level=1)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
        doc.save(str(output))
        metadata = {"path": str(output), "title": title, "paragraphs": len(paragraphs)}
        return ToolResult.ok(
            "docx_generation",
            artifacts=[ToolArtifact(path=str(output), mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", data=metadata)],
            events=[ToolEvent(type="python.docx.generated", payload=metadata)],
            metadata=metadata,
        )
    except Exception as error:  # pragma: no cover - optional dependency/runtime variability
        return ToolResult.fail("docx_generation", str(error), metadata={"path": str(output), "title": title})


TOOL_REGISTRY: dict[str, ToolHandler] = {
    "tool_diagnostics": tool_diagnostics,
    "pdf_probe": pdf_probe,
    "dataset_metadata_probe": dataset_metadata_probe,
    "experiment_log_parse": experiment_log_parse,
    "web_crawl_probe": web_crawl_probe,
    "docx_generation": docx_generation,
}


def call_tool(name: str, input_data: dict[str, Any] | None = None) -> ToolResult:
    handler = TOOL_REGISTRY.get(name)
    if not handler:
        return ToolResult.fail(name, f"Unknown tool: {name}", metadata={"available_tools": sorted(TOOL_REGISTRY)})
    try:
        return handler(input_data or {})
    except Exception as error:
        return ToolResult.fail(name, str(error))

